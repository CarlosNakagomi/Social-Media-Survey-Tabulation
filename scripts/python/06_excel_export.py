"""
06_excel_export.py

Builds the final recruiter-ready Excel tab book for the
W144 Social Media Platform Usage & User Motivations study.

Inputs:
- output/tables/T01.csv ... T28.csv
- tabulation/tabulation_plan.csv
- qa/significance_test_results.csv
- qa/small_base_qa.csv
- qa/qa_summary.csv

Output:
- output/tab_books/Reach3_Social_Media_TabBook.xlsx
"""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter


# ============================================================
# PATHS
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parents[2]

TABLE_DIR = PROJECT_ROOT / "output" / "tables"

PLAN_PATH = (
    PROJECT_ROOT
    / "tabulation"
    / "tabulation_plan.csv"
)

SIG_PATH = (
    PROJECT_ROOT
    / "qa"
    / "significance_test_results.csv"
)

SMALL_BASE_PATH = (
    PROJECT_ROOT
    / "qa"
    / "small_base_qa.csv"
)

QA_SUMMARY_PATH = (
    PROJECT_ROOT
    / "qa"
    / "qa_summary.csv"
)

OUTPUT_DIR = (
    PROJECT_ROOT
    / "output"
    / "tab_books"
)

OUTPUT_DIR.mkdir(
    parents=True,
    exist_ok=True,
)

OUTPUT_PATH = (
    OUTPUT_DIR
    / "Reach3_Social_Media_TabBook.xlsx"
)


# ============================================================
# CONFIGURATION
# ============================================================

PLATFORM_SHEETS = {
    "Facebook": "04_FACEBOOK",
    "Instagram": "05_INSTAGRAM",
    "X": "06_X",
    "TikTok": "07_TIKTOK",
}

CATEGORY_LETTERS = {
    "Age": {
        "18-29": "A",
        "30-49": "B",
        "50-64": "C",
        "65+": "D",
    },
    "Gender": {
        "A man": "A",
        "A woman": "B",
        "In some other way": "C",
    },
    "Education": {
        "College graduate+": "A",
        "Some College": "B",
        "H.S. graduate or less": "C",
    },
    "Region": {
        "Northeast": "A",
        "Midwest": "B",
        "South": "C",
        "West": "D",
    },
    "Income": {
        "Lower income": "A",
        "Middle income": "B",
        "Upper income": "C",
    },
}

RESPONSE_ROWS = [
    "Major reason",
    "Minor reason",
    "Not a reason",
]


# ============================================================
# STYLES
# ============================================================

TITLE_FILL = PatternFill(
    "solid",
    fgColor="1F4E78",
)

SECTION_FILL = PatternFill(
    "solid",
    fgColor="D9EAF7",
)

HEADER_FILL = PatternFill(
    "solid",
    fgColor="5B9BD5",
)

SUBHEADER_FILL = PatternFill(
    "solid",
    fgColor="DDEBF7",
)

BASE_FILL = PatternFill(
    "solid",
    fgColor="F2F2F2",
)

WARNING_FILL = PatternFill(
    "solid",
    fgColor="FFF2CC",
)

VERY_SMALL_FILL = PatternFill(
    "solid",
    fgColor="FCE4D6",
)

WHITE_FONT = Font(
    color="FFFFFF",
    bold=True,
)

TITLE_FONT = Font(
    color="FFFFFF",
    bold=True,
    size=14,
)

HEADER_FONT = Font(
    bold=True,
)

ITALIC_FONT = Font(
    italic=True,
    size=9,
)

THIN_GRAY = Side(
    style="thin",
    color="D9E1F2",
)

BORDER = Border(
    left=THIN_GRAY,
    right=THIN_GRAY,
    top=THIN_GRAY,
    bottom=THIN_GRAY,
)


# ============================================================
# LOAD INPUTS
# ============================================================

plan = pd.read_csv(PLAN_PATH)
sig = pd.read_csv(SIG_PATH)
small_base = pd.read_csv(SMALL_BASE_PATH)
qa_summary = pd.read_csv(QA_SUMMARY_PATH)

assert len(plan) == 28
assert len(sig) == 1596
assert len(small_base) == 476
assert len(qa_summary) == 28

for table_id in plan["table_id"]:
    path = TABLE_DIR / f"{table_id}.csv"

    if not path.exists():
        raise FileNotFoundError(
            f"Missing table export: {path}"
        )


print("=" * 72)
print("W144 EXCEL TAB BOOK EXPORT")
print("=" * 72)

print("\n[PASS] Pipeline inputs loaded")


# ============================================================
# HELPERS
# ============================================================

def parse_column(column_name):

    if column_name == "Total":
        return "Total", "Total"

    if " | " not in column_name:
        raise ValueError(
            f"Invalid table column: {column_name}"
        )

    banner, category = column_name.split(
        " | ",
        1,
    )

    return banner, category


def truthy(value):

    if isinstance(value, bool):
        return value

    return (
        str(value)
        .strip()
        .lower()
        == "true"
    )


def get_significance_letters(
    table_id,
    banner,
    category,
    response,
):

    if banner == "Total":
        return ""

    if (
        banner not in CATEGORY_LETTERS
        or category not in CATEGORY_LETTERS[banner]
    ):
        return ""

    tests = sig.loc[
        (sig["table_id"] == table_id)
        & (sig["banner"] == banner)
        & (sig["response"] == response)
    ]

    letters = []

    for _, test in tests.iterrows():

        if not truthy(
            test["significant_95"]
        ):
            continue

        category_1 = test["category_1"]
        category_2 = test["category_2"]

        pct_1 = float(test["pct_1"])
        pct_2 = float(test["pct_2"])

        if (
            category == category_1
            and pct_1 > pct_2
        ):
            letter = CATEGORY_LETTERS[
                banner
            ].get(category_2)

            if letter:
                letters.append(letter)

        elif (
            category == category_2
            and pct_2 > pct_1
        ):
            letter = CATEGORY_LETTERS[
                banner
            ].get(category_1)

            if letter:
                letters.append(letter)

    order = {
        "A": 1,
        "B": 2,
        "C": 3,
        "D": 4,
    }

    letters = sorted(
        set(letters),
        key=lambda x: order[x],
    )

    return "".join(letters)


def get_display_base(
    table_id,
    column_name,
    raw_base,
):

    base = int(
        round(float(raw_base))
    )

    if column_name == "Total":
        return str(base), None

    match = small_base.loc[
        (small_base["table_id"] == table_id)
        & (small_base["column"] == column_name)
    ]

    if len(match) != 1:
        raise ValueError(
            "Small-base QA lookup failed: "
            f"{table_id} / {column_name}"
        )

    flag = match.iloc[0]["flag"]

    if flag == "VERY SMALL BASE":
        return f"{base}**", flag

    if flag == "SMALL BASE":
        return f"{base}*", flag

    return str(base), flag


def format_percent(value):

    if pd.isna(value):
        return ""

    return f"{float(value):.1f}%"


def style_all_cells(ws):

    for row in ws.iter_rows():
        for cell in row:

            cell.alignment = Alignment(
                vertical="top",
            )

            if cell.value is not None:
                cell.border = BORDER


def set_standard_widths(ws):

    ws.column_dimensions["A"].width = 32

    for column in range(
        2,
        ws.max_column + 1,
    ):
        ws.column_dimensions[
            get_column_letter(column)
        ].width = 16


# ============================================================
# WORKBOOK
# ============================================================

wb = Workbook()

default_sheet = wb.active
wb.remove(default_sheet)


# ============================================================
# 01 README
# ============================================================

ws = wb.create_sheet("01_README")

ws["A1"] = (
    "Social Media Platform Usage & User Motivations"
)

ws["A2"] = (
    "Survey Tabulation, Statistical Testing & QA Workflow"
)

ws["A4"] = "Purpose"

ws["A5"] = (
    "Portfolio project demonstrating an end-to-end survey "
    "tabulation workflow using Pew Research Center American "
    "Trends Panel Wave 144 public-use data."
)

ws["A7"] = "Study Scope"

ws["A8"] = (
    "The analysis compares motivations for using Facebook, "
    "Instagram, X, and TikTok across five demographic banners: "
    "Age, Gender, Education, Region, and Income."
)

ws["A10"] = "Workbook Contents"

ws["A11"] = (
    "28 weighted cross-tabulations, pairwise significance "
    "markers, small-base flags, methodology documentation, "
    "and QA summary."
)

ws["A13"] = "Significance"

ws["A14"] = (
    "Letters indicate statistically significant pairwise "
    "differences within a demographic banner at the 95% "
    "confidence level. A letter identifies the comparison "
    "category with the lower percentage."
)

ws["A16"] = "Small-base notation"

ws["A17"] = (
    "* = unweighted N between 30 and 99; "
    "** = unweighted N below 30."
)

ws["A19"] = "Statistical limitation"

ws["A20"] = (
    "Significance tests use weighted estimates and "
    "Kish-adjusted effective sample sizes. They are "
    "exploratory and are not intended to reproduce Pew "
    "Research Center's exact complex-survey variance "
    "estimation procedure."
)

ws.merge_cells("A1:H1")
ws.merge_cells("A2:H2")

ws["A1"].fill = TITLE_FILL
ws["A1"].font = TITLE_FONT
ws["A1"].alignment = Alignment(
    horizontal="center",
)

ws["A2"].font = Font(
    italic=True,
    size=11,
)

for row in [
    4,
    7,
    10,
    13,
    16,
    19,
]:
    ws[f"A{row}"].font = HEADER_FONT
    ws[f"A{row}"].fill = SECTION_FILL

for row in [
    5,
    8,
    11,
    14,
    17,
    20,
]:

    ws.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=8,
    )

    ws[f"A{row}"].alignment = Alignment(
        wrap_text=True,
        vertical="top",
    )

ws.column_dimensions["A"].width = 30

for column in range(2, 9):
    ws.column_dimensions[
        get_column_letter(column)
    ].width = 14

ws.sheet_view.showGridLines = False


# ============================================================
# 02 CONTENTS
# ============================================================

ws = wb.create_sheet("02_CONTENTS")

ws.append([
    "Sheet",
    "Contents",
])

contents = [
    (
        "01_README",
        "Project overview and interpretation notes",
    ),
    (
        "02_CONTENTS",
        "Workbook navigation",
    ),
    (
        "03_METHODOLOGY",
        "Analytical methodology and QA rules",
    ),
    (
        "04_FACEBOOK",
        "Facebook WHY tables T01-T07",
    ),
    (
        "05_INSTAGRAM",
        "Instagram WHY tables T08-T14",
    ),
    (
        "06_X",
        "X WHY tables T15-T21",
    ),
    (
        "07_TIKTOK",
        "TikTok WHY tables T22-T28",
    ),
    (
        "08_QA_SUMMARY",
        "Table-level QA results",
    ),
]

for item in contents:
    ws.append(item)

for cell in ws[1]:
    cell.fill = TITLE_FILL
    cell.font = WHITE_FONT

ws.column_dimensions["A"].width = 24
ws.column_dimensions["B"].width = 60

ws.freeze_panes = "A2"
ws.sheet_view.showGridLines = False


# ============================================================
# 03 METHODOLOGY
# ============================================================

ws = wb.create_sheet("03_METHODOLOGY")

methodology = [
    (
        "Data source",
        "Pew Research Center American Trends Panel Wave 144 "
        "(March 18-24, 2024).",
    ),
    (
        "Analytical scope",
        "Seven user-motivation items for each of Facebook, "
        "Instagram, X, and TikTok.",
    ),
    (
        "Weighting",
        "Platform-specific W144 survey weights are used for "
        "weighted column percentages.",
    ),
    (
        "Banners",
        "Age, Gender, Education, Census Region, and Income.",
    ),
    (
        "Percentage base",
        "Substantive WHY responses 1, 2, and 3. Special code "
        "99 is excluded from the percentage denominator.",
    ),
    (
        "Effective N",
        "Kish effective sample size: "
        "(sum of weights)^2 / sum of squared weights.",
    ),
    (
        "Significance testing",
        "Two-sided pairwise column-proportion tests using "
        "weighted estimates and Kish-adjusted effective "
        "sample sizes; alpha = 0.05.",
    ),
    (
        "Low-base exclusion",
        "Categories with unweighted N below 100 are excluded "
        "from significance testing.",
    ),
    (
        "Multiplicity",
        "No multiplicity adjustment is applied; significance "
        "markers are exploratory.",
    ),
    (
        "Small-base notation",
        "* indicates N=30-99. ** indicates N<30.",
    ),
    (
        "QA",
        "28 tables, 504 analytical table columns, 1,596 "
        "pairwise tests, and explicit low-base screening.",
    ),
]

ws.append([
    "Methodological Component",
    "Implementation",
])

for item in methodology:
    ws.append(item)

for cell in ws[1]:
    cell.fill = TITLE_FILL
    cell.font = WHITE_FONT

ws.column_dimensions["A"].width = 28
ws.column_dimensions["B"].width = 95

for row in ws.iter_rows():
    for cell in row:

        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

ws.freeze_panes = "A2"
ws.sheet_view.showGridLines = False


# ============================================================
# PLATFORM SHEETS
# ============================================================

for platform, sheet_name in (
    PLATFORM_SHEETS.items()
):

    ws = wb.create_sheet(
        sheet_name
    )

    platform_plan = (
        plan.loc[
            plan["platform"] == platform
        ]
        .copy()
        .sort_values("table_id")
    )

    current_row = 1

    for _, plan_row in (
        platform_plan.iterrows()
    ):

        table_id = plan_row["table_id"]
        measure = plan_row["measure"]

        table_path = (
            TABLE_DIR
            / f"{table_id}.csv"
        )

        table = pd.read_csv(
            table_path,
            index_col=0,
        )

        columns = list(
            table.columns
        )

        # ----------------------------------------------------
        # TABLE TITLE
        # ----------------------------------------------------

        title = (
            f"{table_id} — {platform}: "
            f"{measure}"
        )

        ws.cell(
            current_row,
            1,
            title,
        )

        ws.merge_cells(
            start_row=current_row,
            start_column=1,
            end_row=current_row,
            end_column=len(columns) + 1,
        )

        title_cell = ws.cell(
            current_row,
            1,
        )

        title_cell.fill = TITLE_FILL
        title_cell.font = TITLE_FONT

        title_cell.alignment = Alignment(
            horizontal="left",
        )

        current_row += 1

        # ----------------------------------------------------
        # GROUPED BANNER HEADER
        # ----------------------------------------------------

        header_row = current_row

        ws.cell(
            header_row,
            1,
            "Response",
        )

        ws.cell(
            header_row,
            1,
        ).fill = HEADER_FILL

        ws.cell(
            header_row,
            1,
        ).font = WHITE_FONT

        start_col = 2

        while start_col <= (
            len(columns) + 1
        ):

            column_name = columns[
                start_col - 2
            ]

            banner, _ = parse_column(
                column_name
            )

            end_col = start_col

            while (
                end_col + 1
                <= len(columns) + 1
            ):

                next_column = columns[
                    end_col - 1
                ]

                next_banner, _ = (
                    parse_column(
                        next_column
                    )
                )

                if next_banner != banner:
                    break

                end_col += 1

            if end_col > start_col:

                ws.merge_cells(
                    start_row=header_row,
                    start_column=start_col,
                    end_row=header_row,
                    end_column=end_col,
                )

            cell = ws.cell(
                header_row,
                start_col,
                banner,
            )

            cell.fill = HEADER_FILL
            cell.font = WHITE_FONT

            cell.alignment = Alignment(
                horizontal="center",
            )

            for col in range(
                start_col,
                end_col + 1,
            ):
                ws.cell(
                    header_row,
                    col,
                ).fill = HEADER_FILL

            start_col = end_col + 1

        current_row += 1

        # ----------------------------------------------------
        # CATEGORY HEADER
        # ----------------------------------------------------

        category_row = current_row

        ws.cell(
            category_row,
            1,
            "",
        )

        for col_idx, column_name in enumerate(
            columns,
            start=2,
        ):

            banner, category = (
                parse_column(
                    column_name
                )
            )

            if banner == "Total":
                label = "Total"

            else:
                letter = CATEGORY_LETTERS[
                    banner
                ][category]

                label = (
                    f"{letter}. {category}"
                )

            cell = ws.cell(
                category_row,
                col_idx,
                label,
            )

            cell.fill = SUBHEADER_FILL
            cell.font = HEADER_FONT

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        current_row += 1

        # ----------------------------------------------------
        # BASE ROW
        # ----------------------------------------------------

        base_row = current_row

        ws.cell(
            base_row,
            1,
            "Unweighted Base",
        )

        ws.cell(
            base_row,
            1,
        ).fill = BASE_FILL

        ws.cell(
            base_row,
            1,
        ).font = HEADER_FONT

        for col_idx, column_name in enumerate(
            columns,
            start=2,
        ):

            raw_base = table.loc[
                "Unweighted Base",
                column_name,
            ]

            display_base, flag = (
                get_display_base(
                    table_id,
                    column_name,
                    raw_base,
                )
            )

            cell = ws.cell(
                base_row,
                col_idx,
                display_base,
            )

            cell.alignment = Alignment(
                horizontal="center",
            )

            if flag == "SMALL BASE":
                cell.fill = WARNING_FILL

            elif flag == "VERY SMALL BASE":
                cell.fill = VERY_SMALL_FILL

            else:
                cell.fill = BASE_FILL

        current_row += 1

        # ----------------------------------------------------
        # RESPONSE ROWS
        # ----------------------------------------------------

        for response in RESPONSE_ROWS:

            ws.cell(
                current_row,
                1,
                response,
            )

            for col_idx, column_name in enumerate(
                columns,
                start=2,
            ):

                banner, category = (
                    parse_column(
                        column_name
                    )
                )

                value = table.loc[
                    response,
                    column_name,
                ]

                display_value = (
                    format_percent(
                        value
                    )
                )

                letters = (
                    get_significance_letters(
                        table_id,
                        banner,
                        category,
                        response,
                    )
                )

                if letters:
                    display_value = (
                        f"{display_value} "
                        f"{letters}"
                    )

                cell = ws.cell(
                    current_row,
                    col_idx,
                    display_value,
                )

                cell.alignment = Alignment(
                    horizontal="center",
                )

            current_row += 1

        # ----------------------------------------------------
        # FOOTNOTE
        # ----------------------------------------------------

        footnote = (
            "Letters indicate significantly higher percentages "
            "than the referenced category within the same banner "
            "(two-sided 95% test). "
            "* N=30-99; ** N<30. "
            "Categories with N<100 are excluded from significance "
            "testing. No multiplicity adjustment."
        )

        ws.cell(
            current_row,
            1,
            footnote,
        )

        ws.merge_cells(
            start_row=current_row,
            start_column=1,
            end_row=current_row,
            end_column=len(columns) + 1,
        )

        ws.cell(
            current_row,
            1,
        ).font = ITALIC_FONT

        ws.cell(
            current_row,
            1,
        ).alignment = Alignment(
            wrap_text=True,
        )

        current_row += 3

    set_standard_widths(ws)
    style_all_cells(ws)

    ws.freeze_panes = "B4"
    ws.sheet_view.showGridLines = False


# ============================================================
# 08 QA SUMMARY
# ============================================================

ws = wb.create_sheet(
    "08_QA_SUMMARY"
)

qa_display = (
    qa_summary.copy()
)

for row_idx, values in enumerate(
    qa_display.itertuples(
        index=False,
        name=None,
    ),
    start=2,
):

    for col_idx, value in enumerate(
        values,
        start=1,
    ):

        ws.cell(
            row_idx,
            col_idx,
            value,
        )


for col_idx, column in enumerate(
    qa_display.columns,
    start=1,
):

    cell = ws.cell(
        1,
        col_idx,
        column,
    )

    cell.fill = TITLE_FILL
    cell.font = WHITE_FONT

    cell.alignment = Alignment(
        horizontal="center",
        wrap_text=True,
    )


for row in ws.iter_rows(
    min_row=2,
):

    for cell in row:

        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )


for col_idx, column in enumerate(
    qa_display.columns,
    start=1,
):

    if column in {
        "variable",
        "weight",
        "measure",
        "universe",
    }:
        width = 38

    else:
        width = 18

    ws.column_dimensions[
        get_column_letter(
            col_idx
        )
    ].width = width


ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
ws.sheet_view.showGridLines = False


# ============================================================
# SAVE
# ============================================================

wb.save(
    OUTPUT_PATH
)

assert OUTPUT_PATH.exists()

print(
    "[PASS] Excel workbook generated:"
)

print(
    f"       {OUTPUT_PATH}"
)


# ============================================================
# ROUND-TRIP WORKBOOK QA
# ============================================================

check_wb = load_workbook(
    OUTPUT_PATH,
    read_only=True,
    data_only=False,
)

expected_sheets = [
    "01_README",
    "02_CONTENTS",
    "03_METHODOLOGY",
    "04_FACEBOOK",
    "05_INSTAGRAM",
    "06_X",
    "07_TIKTOK",
    "08_QA_SUMMARY",
]

assert (
    check_wb.sheetnames
    == expected_sheets
), (
    "Unexpected workbook sheets: "
    f"{check_wb.sheetnames}"
)

assert (
    check_wb[
        "04_FACEBOOK"
    ]["A1"].value.startswith(
        "T01"
    )
)

assert (
    check_wb[
        "05_INSTAGRAM"
    ]["A1"].value.startswith(
        "T08"
    )
)

assert (
    check_wb[
        "06_X"
    ]["A1"].value.startswith(
        "T15"
    )
)

assert (
    check_wb[
        "07_TIKTOK"
    ]["A1"].value.startswith(
        "T22"
    )
)

assert (
    check_wb[
        "08_QA_SUMMARY"
    ].max_row
    == 29
)

check_wb.close()

print(
    "[PASS] Workbook structure validated"
)


# ============================================================
# SIGNIFICANCE MARKER QA
# ============================================================

# Build the exact set of analytical cells that should receive
# at least one significance letter.

significant_tests = sig.loc[
    sig[
        "significant_95"
    ].apply(truthy)
].copy()

expected_marker_cells = set()

for _, test in (
    significant_tests.iterrows()
):

    table_id = test["table_id"]
    banner = test["banner"]
    response = test["response"]

    category_1 = test["category_1"]
    category_2 = test["category_2"]

    pct_1 = float(
        test["pct_1"]
    )

    pct_2 = float(
        test["pct_2"]
    )

    if pct_1 > pct_2:
        category = category_1

    elif pct_2 > pct_1:
        category = category_2

    else:
        continue

    expected_marker_cells.add(
        (
            table_id,
            banner,
            response,
            category,
        )
    )


assert (
    len(expected_marker_cells)
    == 342
), (
    "Expected 342 unique "
    "significance-marker cells; "
    f"found "
    f"{len(expected_marker_cells)}"
)


# Independently determine which analytical cells the
# workbook-generation logic assigns letters to.

actual_marker_cells = set()

for _, plan_row in (
    plan.iterrows()
):

    table_id = (
        plan_row["table_id"]
    )

    table_path = (
        TABLE_DIR
        / f"{table_id}.csv"
    )

    table = pd.read_csv(
        table_path,
        index_col=0,
    )

    for response in RESPONSE_ROWS:

        for column_name in (
            table.columns
        ):

            banner, category = (
                parse_column(
                    column_name
                )
            )

            letters = (
                get_significance_letters(
                    table_id,
                    banner,
                    category,
                    response,
                )
            )

            if letters:

                actual_marker_cells.add(
                    (
                        table_id,
                        banner,
                        response,
                        category,
                    )
                )


assert (
    actual_marker_cells
    == expected_marker_cells
), (
    "Excel significance-marker "
    "cells do not match the "
    "significance-test results."
)

marker_cells = len(
    actual_marker_cells
)

assert marker_cells == 342

print(
    "[PASS] Significance markers "
    "reconciled exactly "
    f"({marker_cells}/342 cells)"
)


# ============================================================
# SMALL-BASE OUTPUT QA
# ============================================================

small_count = int(
    (
        small_base["flag"]
        == "SMALL BASE"
    ).sum()
)

very_small_count = int(
    (
        small_base["flag"]
        == "VERY SMALL BASE"
    ).sum()
)

ok_count = int(
    (
        small_base["flag"]
        == "OK"
    ).sum()
)

assert ok_count == 448
assert small_count == 21
assert very_small_count == 7

print(
    "[PASS] Small-base distribution "
    "reconciled "
    "(448 OK | 21 SMALL | 7 VERY SMALL)"
)


# ============================================================
# FINAL STATUS
# ============================================================

print(
    "\n"
    + "-"
    * 72
)

print(
    "EXCEL EXPORT SUMMARY"
)

print(
    "-"
    * 72
)

print(
    "Workbook sheets:        8"
)

print(
    "Platform tables:        28"
)

print(
    f"Significance markers:   "
    f"{marker_cells} cells"
)

print(
    "Small-base notation:    "
    "* / ** applied"
)

print(
    "\n"
    + "="
    * 72
)

print(
    "EXCEL TAB BOOK EXPORT PASSED"
)

print(
    "="
    * 72
)